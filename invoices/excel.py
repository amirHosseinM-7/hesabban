"""Excel exports (openpyxl)."""
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from buildings.jalali import format_jalali

HEADER_FILL = PatternFill("solid", fgColor="0A84FF")
HEADER_FONT = Font(name="Vazirmatn", bold=True, color="FFFFFF")
BASE_FONT = Font(name="Vazirmatn")


def _workbook(title, headers):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.sheet_view.rightToLeft = True
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    return wb, ws


def _finish(wb, ws, response, count):
    for row in range(2, count + 2):
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).font = BASE_FONT
    widths = [18, 22, 26, 16, 30]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width
    wb.save(response)


def export_payments(building, response):
    from payments.models import Payment

    wb, ws = _workbook("پرداخت‌ها", ["واحد", "مبلغ (تومان)", "تاریخ", "روش", "یادداشت"])
    payments = Payment.objects.filter(unit__building=building).select_related("unit").order_by("-date", "-id")
    row = 2
    for p in payments:
        ws.cell(row=row, column=1, value=f"واحد {p.unit.number}")
        ws.cell(row=row, column=2, value=float(p.amount))
        ws.cell(row=row, column=3, value=format_jalali(p.date))
        ws.cell(row=row, column=4, value=p.get_method_display())
        ws.cell(row=row, column=5, value=p.note)
        row += 1
    ws.cell(row=row, column=1, value="جمع")
    ws.cell(row=row, column=2, value=float(sum((p.amount for p in payments), Decimal("0"))))
    ws.cell(row=row, column=1).font = Font(name="Vazirmatn", bold=True)
    ws.cell(row=row, column=2).font = Font(name="Vazirmatn", bold=True)
    _finish(wb, ws, response, payments.count())
    return response


def export_expenses(building, response):
    wb, ws = _workbook("هزینه‌ها", ["عنوان", "دسته", "مبلغ (تومان)", "تاریخ", "یادداشت"])
    expenses = building.expenses.order_by("-date", "-id")
    row = 2
    for e in expenses:
        ws.cell(row=row, column=1, value=e.title)
        ws.cell(row=row, column=2, value=e.get_category_display())
        ws.cell(row=row, column=3, value=float(e.amount))
        ws.cell(row=row, column=4, value=format_jalali(e.date))
        ws.cell(row=row, column=5, value=e.note)
        row += 1
    ws.cell(row=row, column=1, value="جمع")
    ws.cell(row=row, column=3, value=float(sum((e.amount for e in expenses), Decimal("0"))))
    ws.cell(row=row, column=1).font = Font(name="Vazirmatn", bold=True)
    ws.cell(row=row, column=3).font = Font(name="Vazirmatn", bold=True)
    _finish(wb, ws, response, expenses.count())
    return response
